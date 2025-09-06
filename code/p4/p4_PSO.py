# -*- coding: utf-8 -*-
"""
Problem 4: 三架无人机各投放一枚烟幕弹，对单枚导弹 M1 的协同遮蔽优化（PSO）
依赖: numpy, matplotlib, openpyxl

输出:
- code/p4/output_PSO/summary_p4_PSO.txt
- code/p4/output_PSO/pso_progress.png
- 附件/result2.xlsx（若存在则更新首个工作表，若不存在则创建）
"""
import os
import time
import math
import random
import numpy as np
import matplotlib.pyplot as plt

try:
    from openpyxl import load_workbook, Workbook
except Exception:
    load_workbook = None
    Workbook = None

# Matplotlib 中文显示
plt.rcParams['font.sans-serif'] = ['SimHei']
plt.rcParams['axes.unicode_minus'] = False

# -----------------------------
# 1) 物理参数与初始条件
# -----------------------------
g = 9.80665
SMOKE_EFFECTIVE_RADIUS = 10.0
SINK_SPEED = 3.0
EFFECTIVE_TIME = 20.0
MISSILE_SPEED = 300.0

# 三架无人机初始位置（FY1, FY2, FY3）
UAV_STARTS = [
    np.array([17800.0,     0.0, 1800.0], dtype=float),  # FY1
    np.array([12000.0,  1400.0, 1400.0], dtype=float),  # FY2
    np.array([ 6000.0, -3000.0,  700.0], dtype=float),  # FY3
]

# 导弹与目标
MISSILE_START_M1 = np.array([20000.0, 0.0, 2000.0], dtype=float)
TARGET_CENTER = np.array([0.0, 200.0, 5.0], dtype=float)

R = SMOKE_EFFECTIVE_RADIUS
cloud_sink = SINK_SPEED
M0 = MISSILE_START_M1
to_origin = -M0 / np.linalg.norm(M0)
vM = MISSILE_SPEED * to_origin
T_center = TARGET_CENTER

def M_pos(t: float) -> np.ndarray:
    return M0 + vM * t

# -----------------------------
# 2) 几何与适应度
# -----------------------------
def point_to_segment_distance(c: np.ndarray, a: np.ndarray, b: np.ndarray):
    ab = b - a
    denom = float(np.dot(ab, ab))
    if denom < 1e-9:
        return np.linalg.norm(c - a), 0.0
    lam = float(np.dot(c - a, ab) / denom)
    lam_clamped = max(0.0, min(1.0, lam))
    d = float(np.linalg.norm(c - (a + lam_clamped * ab)))
    return d, lam_clamped

def _unpack_params(params):
    """
    params 维度 = 12
    对每个 UAV i∈{0,1,2}： [angle_i, speed_i, t_release_i, t_fuze_i]
    angle 单位: 度, speed 单位: m/s, 时间单位: s
    """
    triples = []
    for i in range(3):
        angle = params[4*i + 0]
        speed = params[4*i + 1]
        t_rel = params[4*i + 2]
        t_fuz = params[4*i + 3]
        triples.append((angle, speed, t_rel, t_fuz))
    return triples

def calculate_total_coverage_union(params, N_grid=1001, is_final_calc=False):
    # 钳制边界
    clamped = [check_bounds(v, PARAM_BOUNDS[i]) for i, v in enumerate(params)]
    triples = _unpack_params(clamped)

    bombs_data = []
    for i, (angle, speed, t_rel, t_fuz) in enumerate(triples):
        angle_rad = np.deg2rad(angle % 360.0)
        uav_dir = np.array([np.cos(angle_rad), np.sin(angle_rad), 0.0])
        U0 = UAV_STARTS[i]
        # 投放点（等高直线飞行）
        p_release = U0 + uav_dir * speed * t_rel
        # 起爆点（释放后：水平匀速、竖直自由落体，仅考虑引信延时 t_fuz）
        explosion_point = np.array([
            p_release[0] + uav_dir[0] * speed * t_fuz,
            p_release[1] + uav_dir[1] * speed * t_fuz,
            U0[2] - 0.5 * g * (t_fuz ** 2)
        ])
        t_explode = t_rel + t_fuz
        bombs_data.append({'E': explosion_point, 't_exp': t_explode, 'P_rel': p_release, 'uav_idx': i})

    min_exp_time = min(b['t_exp'] for b in bombs_data)
    t0 = min_exp_time
    t1 = max(b['t_exp'] + EFFECTIVE_TIME for b in bombs_data)
    t_grid = np.linspace(t0, t1, N_grid)
    dt = t_grid[1] - t_grid[0]

    union_mask = np.zeros_like(t_grid, dtype=bool)
    for idx, t in enumerate(t_grid):
        m_pos = M_pos(t)
        flag = False
        for bomb in bombs_data:
            if t >= bomb['t_exp'] and t <= bomb['t_exp'] + EFFECTIVE_TIME:
                c_pos = bomb['E'] - np.array([0, 0, cloud_sink * (t - bomb['t_exp'])])
                d, _ = point_to_segment_distance(c_pos, m_pos, T_center)
                if d <= R:
                    flag = True
                    break
        union_mask[idx] = flag

    total_duration = float(np.sum(union_mask) * dt)
    if is_final_calc:
        return total_duration, bombs_data
    return total_duration

# 缓存
FITNESS_CACHE = {}

def get_fitness(params, fidelity='low'):
    key = (tuple([round(check_bounds(v, PARAM_BOUNDS[i]), 3) for i, v in enumerate(params)]), fidelity)
    if key in FITNESS_CACHE:
        return FITNESS_CACHE[key]
    N = 301 if fidelity == 'low' else 1501
    val = calculate_total_coverage_union(params, N_grid=N, is_final_calc=False)
    FITNESS_CACHE[key] = val
    return val

# -----------------------------
# 3) PSO
# -----------------------------
# 每个 UAV: 角度、速度、释放时间、引信时间
PARAM_BOUNDS = [
    (0.0, 360.0), (70.0, 140.0), (0.0, 15.0), (0.005, 12.0),  # FY1
    (0.0, 360.0), (70.0, 140.0), (0.0, 15.0), (0.005, 12.0),  # FY2
    (0.0, 360.0), (70.0, 140.0), (0.0, 15.0), (0.005, 12.0),  # FY3
]

def check_bounds(val, bounds):
    return max(bounds[0], min(bounds[1], float(val)))

def project_params(params):
    q = [check_bounds(v, b) for v, b in zip(params, PARAM_BOUNDS)]
    # 角度归一
    for i in [0, 4, 8]:
        q[i] = q[i] % 360.0
    return q

def run_pso_optimization():
    print("=== 开始执行三机一弹协同策略优化 (PSO, Q4) ===")
    random.seed(42); np.random.seed(42)

    SWARM_SIZE = 90
    ITERATIONS = 500
    W = 0.6
    C1 = 1.6
    C2 = 1.6

    dim = len(PARAM_BOUNDS)
    positions = np.array([[random.uniform(b[0], b[1]) for b in PARAM_BOUNDS] for _ in range(SWARM_SIZE)], dtype=float)
    velocities = np.zeros((SWARM_SIZE, dim), dtype=float)

    pbest = positions.copy()
    pbest_fit = np.array([get_fitness(project_params(list(p)), fidelity='low') for p in pbest])
    g_idx = int(np.argmax(pbest_fit))
    gbest = pbest[g_idx].copy()
    gbest_fit = float(pbest_fit[g_idx])

    start_time = time.time()
    gen_best_list = []

    for it in range(ITERATIONS):
        r1 = np.random.rand(SWARM_SIZE, dim)
        r2 = np.random.rand(SWARM_SIZE, dim)
        velocities = W*velocities + C1*r1*(pbest - positions) + C2*r2*(gbest - positions)
        positions = positions + velocities
        positions = np.array([project_params(list(p)) for p in positions], dtype=float)

        fits = np.array([get_fitness(list(p), fidelity='low') for p in positions])
        gen_best_list.append(float(np.max(fits)))

        improved = fits > pbest_fit
        pbest[improved] = positions[improved]
        pbest_fit[improved] = fits[improved]

        g_idx = int(np.argmax(pbest_fit))
        gbest_candidate = pbest[g_idx].copy()
        gbest_candidate_fit = get_fitness(list(gbest_candidate), fidelity='high')
        if gbest_candidate_fit > gbest_fit:
            gbest, gbest_fit = gbest_candidate, gbest_candidate_fit
            print("[刷新历史遮挡最长] 当前最佳(高保真): {:.3f} s".format(gbest_fit))

        print(f"迭代 {it+1:3d}/{ITERATIONS} | 本轮最佳(低保真): {np.max(fits):.3f} s | 历史最佳(高保真): {gbest_fit:.3f} s")

    # 输出摘要与进度图
    out_dir = os.path.join(os.path.dirname(__file__), 'output_PSO')
    os.makedirs(out_dir, exist_ok=True)
    try:
        with open(os.path.join(out_dir, 'summary_p4_PSO.txt'), 'w', encoding='utf-8') as f:
            f.write("[总体最优策略(高保真评估)]\n")
            f.write("最长有效遮蔽并集时长: {:.3f} s\n".format(gbest_fit))
            triples = _unpack_params(gbest)
            for i, (angle, speed, t_rel, t_fuz) in enumerate(triples):
                f.write(f"FY{i+1}: 角度 {angle:.3f} 度, 速度 {speed:.2f} m/s, 投放 {t_rel:.3f} s, 引信 {t_fuz:.3f} s\n")
    except Exception as e:
        print("写入摘要失败:", e)

    try:
        plt.figure(figsize=(10,4.5))
        x = list(range(1, len(gen_best_list)+1))
        plt.plot(x, gen_best_list, 'o-', linewidth=1.8, markersize=4, color='#1f77b4', alpha=0.9, label='本代最长', zorder=2)
        plt.xlabel('代数'); plt.ylabel('遮蔽时间 (s)'); plt.title('PSO 优化进度 (Q4)')
        plt.legend(); plt.grid(alpha=0.3); plt.tight_layout()
        plt.savefig(os.path.join(out_dir, 'pso_progress.png'), dpi=150)
        plt.close()
    except Exception as e:
        print("保存进度图失败:", e)

    return list(gbest), gbest_fit

# -----------------------------
# 4) 结果输出（控制台与 Excel）
# -----------------------------
def calculate_individual_durations(bombs_data, N_grid=4001):
    min_exp_time = min(b['t_exp'] for b in bombs_data)
    t1 = max(b['t_exp'] + EFFECTIVE_TIME for b in bombs_data)
    t_grid = np.linspace(min_exp_time, t1, N_grid)
    dt = t_grid[1] - t_grid[0]
    durations = []
    for bomb in bombs_data:
        mask = np.zeros_like(t_grid, dtype=bool)
        for i, t in enumerate(t_grid):
            if t >= bomb['t_exp'] and t <= bomb['t_exp'] + EFFECTIVE_TIME:
                c_pos = bomb['E'] - np.array([0, 0, cloud_sink * (t - bomb['t_exp'])])
                d, _ = point_to_segment_distance(c_pos, M_pos(t), T_center)
                if d <= R:
                    mask[i] = True
        durations.append(float(np.sum(mask) * dt))
    return durations

def format_and_output_results(best_params):
    print("\n=== 对最优策略进行高精度计算并生成报告 (Q4) ===")
    total_duration_precise, bombs_data_precise = calculate_total_coverage_union(
        best_params, N_grid=5001, is_final_calc=True
    )
    triples = _unpack_params(best_params)

    print("总有效遮蔽并集时长 (s): {:.4f}".format(total_duration_precise))
    for i, bomb in enumerate(bombs_data_precise):
        pr = bomb['P_rel']; e = bomb['E']
        angle, speed, t_rel, t_fuz = triples[i]
        print(f"FY{i+1}")
        print("  航向角(度): {:.2f}".format(angle))
        print("  速度(m/s): {:.2f}".format(speed))
        print("  投放时间(s): {:.3f}".format(t_rel))
        print("  引信时间(s): {:.3f}".format(t_fuz))
        print("  投放点(m): ({:.3f}, {:.3f}, {:.3f})".format(pr[0], pr[1], pr[2]))
        print("  起爆点(m): ({:.3f}, {:.3f}, {:.3f})".format(e[0], e[1], e[2]))

    indiv = calculate_individual_durations(bombs_data_precise, N_grid=5001)
    for i in range(3):
        print("FY{} 独立有效遮蔽时长(s): {:.3f}".format(i+1, indiv[i]))

    # 写入 Excel: 附件/result2.xlsx
    if Workbook is None:
        print("未安装 openpyxl，跳过 Excel 写入。请在 requirements.txt 中包含 openpyxl 并安装。")
        return

    proj_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
    xlsx_path = os.path.join(proj_root, '附件', 'result2.xlsx')
    os.makedirs(os.path.dirname(xlsx_path), exist_ok=True)
    wb = None
    if os.path.exists(xlsx_path) and load_workbook is not None:
        try:
            wb = load_workbook(xlsx_path)
        except Exception:
            wb = Workbook()
    else:
        wb = Workbook()

    ws = wb.active
    ws.title = ws.title or 'Sheet1'

    # 写入表头（若空表）
    if ws.max_row == 1 and ws.max_column == 1 and ws['A1'].value is None:
        ws.append([
            'UAV', '航向角(度)', '速度(m/s)', '投放时间(s)', '引信时间(s)',
            '投放点x(m)', '投放点y(m)', '投放点z(m)',
            '起爆点x(m)', '起爆点y(m)', '起爆点z(m)'
        ])

    # 清理旧数据（可选：这里简单追加）
    for i, bomb in enumerate(bombs_data_precise):
        pr = bomb['P_rel']; e = bomb['E']
        angle, speed, t_rel, t_fuz = triples[i]
        ws.append([
            f'FY{i+1}', round(angle, 3), round(speed, 3), round(t_rel, 3), round(t_fuz, 3),
            round(pr[0], 3), round(pr[1], 3), round(pr[2], 3),
            round(e[0], 3), round(e[1], 3), round(e[2], 3)
        ])

    # 总体遮蔽时长写入最后一行下方
    ws.append(['总有效遮蔽并集时长(s)', round(total_duration_precise, 3)])

    try:
        wb.save(xlsx_path)
        print("Excel 已写入:", xlsx_path)
    except Exception as e:
        print("保存 Excel 失败:", e)


if __name__ == "__main__":
    best_params, best_fit = run_pso_optimization()
    if best_params:
        format_and_output_results(best_params)


